from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = '17365D'
BLUE = 'D9EAF7'
LIGHT = 'F4F7FA'
GRID = 'D9E2F3'
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        props.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_borders(cell, color=GRID, size='6'):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        props.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        props.append(margins)
    for side, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = margins.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = field
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, text, end])


def clear_paragraph_borders(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn('w:pBdr'))
    if borders is not None:
        ppr.remove(borders)


def clear_style_borders(style):
    ppr = style._element.pPr
    if ppr is not None:
        borders = ppr.find(qn('w:pBdr'))
        if borders is not None:
            ppr.remove(borders)


def add_inline(paragraph, text):
    text = text.replace('`', '')
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_body(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, text)
    return paragraph


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(rows[0]):
        cell = header.cells[i]
        cell.text = value
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
    for row_index, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.text = value
            set_cell_shading(cell, LIGHT if row_index % 2 else 'FFFFFF')
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def build(source: Path, output: Path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [('Title', 28, BLACK), ('Heading 1', 18, BLACK), ('Heading 2', 13, BLACK), ('Heading 3', 11, BLACK)]:
        style = styles[name]
        style.font.name = 'Aptos Display' if name in ('Title', 'Heading 1') else 'Aptos'
        style._element.rPr.rFonts.set(qn('w:ascii'), style.font.name)
        style._element.rPr.rFonts.set(qn('w:hAnsi'), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        clear_style_borders(style)

    # Header/footer: quiet, functional, and consistent.
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run('PROPBREW  |  MASTER PRODUCT REQUIREMENTS DOCUMENT')
    run.font.name = 'Aptos'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Propbrew confidential  |  Page ')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)
    add_field(footer, 'PAGE')

    # Cover page.
    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run('Propbrew Master Product Requirements Document')
    clear_paragraph_borders(title)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run('Real Estate Operating System for Nigerian Estate Operators')
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(68, 68, 68)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_repeat_table_header(meta.rows[0])
    metadata = [
        ('Pilot market', 'Lagos, Nigeria'),
        ('Document version', '1.0'),
        ('Date', '11 October 2026'),
        ('Status', 'Product baseline; production launch gated by infrastructure, UAT, and client sign-off'),
    ]
    for row, (label, value) in zip(meta.rows, metadata):
        row.cells[0].text = label
        row.cells[1].text = value
        for i, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_margins(cell, top=140, bottom=140)
            set_cell_shading(cell, BLUE if i == 0 else 'FFFFFF')
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for r in paragraph.runs:
                    r.font.size = Pt(10)
                    if i == 0:
                        r.bold = True
    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(12)
    intro.add_run('Purpose. ').bold = True
    intro.add_run('This document is the operating baseline for product, engineering, operations, finance, security, and client sign-off. It defines what Propbrew must do, how the Lagos pilot is measured, and which conditions must be true before production promotion.')
    doc.add_page_break()

    lines = source.read_text(encoding='utf-8').splitlines()
    i = 0
    source_content_started = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith('# '):
            i += 1
            continue  # Cover title already rendered.
        if not source_content_started:
            if line.startswith('## '):
                source_content_started = True
            else:
                i += 1
                continue  # Cover metadata already rendered.
        if line.startswith('## '):
            add_body(doc, line[3:].strip(), 'Heading 1')
            i += 1
            continue
        if line.startswith('### '):
            add_body(doc, line[4:].strip(), 'Heading 2')
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                raw = lines[i].strip().strip('|')
                cells = [c.strip() for c in raw.split('|')]
                if not all(re.fullmatch(r':?-{3,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            if len(rows) >= 2:
                add_table(doc, rows)
            continue
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, line[2:].strip())
            i += 1
            continue
        if re.match(r'^\d+\. ', line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            add_inline(p, line)
            i += 1
            continue
        add_body(doc, line)
        i += 1

    # Keep document properties neutral and professional.
    doc.core_properties.title = 'Propbrew Master Product Requirements Document'
    doc.core_properties.subject = 'Product requirements and launch acceptance baseline'
    doc.core_properties.author = 'Propbrew Product Team'
    doc.core_properties.keywords = 'Propbrew, PRD, real estate operations, Nigeria, product requirements'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit('Usage: build_master_prd_docx.py SOURCE.md OUTPUT.docx')
    build(Path(sys.argv[1]), Path(sys.argv[2]))
